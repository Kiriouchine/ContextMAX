"""Workbooks as values: each number with its label, unit, column title, address and formula.

Read row by row, a sheet loses which label a number belongs to, its unit and, worst, its
formula, which may be the only place a design input is written down. So every value becomes
one line of the text cache,

    <label> = <value> <unit>  @<sheet>!<cell>  [column <title>]  {<formula>}

which a plain search finds by label, number, unit or formula constant, and which can be cited
by its address. The same facts go to the sheet's cell file, one JSON record per cell; that file
feeds the value list.

  .xlsx .xlsm  openpyxl, opened twice: formulas from one pass, the results Excel stored from the other
  .xls         xlrd for the values; formulas from an .xlsx copy made by Excel, whose values are never used
  .csv         the standard library, with the same row limits
"""
from __future__ import annotations

import codecs
import csv
import hashlib
import json
import math
import os
import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import NamedTuple

from contextmax.docs.readers import win_path, write_out

warnings.filterwarnings("ignore", module="openpyxl")

DATA_TABLE_ROWS = 400        # a longer sheet is a data table: listed in part
DATA_TABLE_SAMPLE = 60
CELLS_PER_SHEET = 400_000
LEFT_LOOK = 4                # how far to the left a label may stand
UP_LOOK = 12                 # how far up a column title may stand
SHEET_SCHEMA_VERSION = 1
SHEET_READER_VERSION = 2

# Unit spellings folded to one form; a project adds its own under vocabulary.units.
# A unit not listed here is kept as written, in brackets.
_UNIT_GROUPS = (
    ("deg", ("deg", "degree", "degrees", "°")), ("rad", ("rad",)), ("rad/s", ("rad/s",)), ("rpm", ("rpm",)),
    ("counts", ("counts", "count", "pulses", "inc", "incr", "increments")),
    ("mm", ("mm",)), ("cm", ("cm",)), ("m", ("m",)), ("km", ("km",)), ("um", ("um", "µm")), ("nm", ("nm",)),
    ("kg", ("kg",)), ("g", ("g", "gram")), ("mg", ("mg",)), ("N", ("n",)), ("kN", ("kn",)), ("mNm", ("mnm",)),
    ("s", ("s", "sec")), ("ms", ("ms",)), ("us", ("us", "µs")), ("min", ("min",)), ("h", ("h",)),
    ("Hz", ("hz",)), ("kHz", ("khz",)), ("MHz", ("mhz",)), ("GHz", ("ghz",)),
    ("V", ("v",)), ("mV", ("mv",)), ("kV", ("kv",)), ("A", ("a",)), ("mA", ("ma",)), ("uA", ("ua",)),
    ("W", ("w",)), ("mW", ("mw",)), ("kW", ("kw",)), ("%", ("%",)), ("ppm", ("ppm",)), ("dB", ("db",)),
    ("bar", ("bar",)), ("mbar", ("mbar",)), ("Pa", ("pa",)), ("kPa", ("kpa",)), ("MPa", ("mpa",)),
    ("K", ("k",)), ("°C", ("°c", "c")), ("ohm", ("ohm", "ω")), ("kohm", ("kohm",)),
    ("F", ("f",)), ("uF", ("uf",)), ("nF", ("nf",)), ("pF", ("pf",)), ("J", ("j",)), ("mJ", ("mj",)),
    ("m/s", ("m/s",)), ("mm/s", ("mm/s",)), ("m/s2", ("m/s2", "m/s^2")), ("mm/s2", ("mm/s2",)),
    ("bit", ("bit",)), ("byte", ("byte",)), ("kB", ("kb",)), ("MB", ("mb",)),
)
DEFAULT_UNITS = {spelling: unit for unit, spellings in _UNIT_GROUPS for spelling in spellings}
# Spellings where case decides: megapascal against millipascal, newton-metre against nanometre,
# megabit against millibar.
CASE_KEPT = {"MPa": "MPa", "Mb": "Mbit", "Nm": "Nm", "mNm": "mNm", "mPa": "mPa", "mb": "mbar", "nm": "nm"}
NO_UNIT = frozenset({"n/a", "n.a.", "na", "[-]", "--", "-", ""})

_LOWER_WORDS = re.compile(r"[a-zà-ÿ]+")
_BRACKETED = re.compile(r"\[.{1,16}\]")
_DECIMAL_ZEROS = re.compile(r"[.](0+)")
_DATES = (datetime, date, time)
_KINDS = ((bool, "boolean"), ((int, float), "number"), (_DATES, "datetime"))     # bool first: True is an int
_XLS_SHEET_STATES = {0: "visible", 1: "hidden", 2: "veryHidden"}


def kind_of(value) -> str:
    if value is None:
        return "empty"
    return next((kind for types, kind in _KINDS if isinstance(value, types)), "text")


@dataclass(slots=True)
class SheetCell:
    value: object = None
    formula: str | None = None
    fmt: object = None
    hidden: bool = False
    kind: str | None = field(default=None)

    def __post_init__(self):
        self.kind = self.kind or kind_of(self.value)


class Sheet(NamedTuple):
    name: str
    grid: dict                   # (row, column), both from 1 -> SheetCell
    rows: int                    # rows read; fewer than the sheet has when `note` says so
    cols: int
    state: str                   # visible, hidden or veryHidden
    note: str


class Units:
    """The unit spellings of one project: the defaults plus what project.json adds."""

    def __init__(self, extra: dict | None = None):
        self.table = {**DEFAULT_UNITS, **{str(k).lower(): v for k, v in (extra or {}).items()}}

    def read(self, token) -> tuple[str, bool]:
        """(the unit, whether it is a known one). Unknown units come back in brackets."""
        if token is None:
            return "", False
        text = str(token).strip()
        inner = text[1:-1].strip() if text.startswith("[") and text.endswith("]") else text
        if text.lower() in NO_UNIT or inner.lower() in NO_UNIT:
            return "", True
        known = CASE_KEPT.get(inner) or self.table.get(inner.lower())
        return (known, True) if known else (f"[{inner}]", False)

    def is_unit(self, text) -> bool:
        if text is None:
            return False
        t = str(text).strip()
        return 0 < len(t) <= 18 and (bool(_BRACKETED.fullmatch(t)) or t.lower() in self.table or t.lower() in NO_UNIT)


def column_name(n: int) -> str:
    """1 -> A, 27 -> AA."""
    letters = []
    while n > 0:
        n, rest = divmod(n - 1, 26)
        letters.append(chr(ord("A") + rest))
    return "".join(reversed(letters))


def _trim(x: float) -> str:
    return repr(x).rstrip("0").rstrip(".")


def _decimals(number_format: str) -> int | None:
    m = _DECIMAL_ZEROS.search(number_format)
    return len(m.group(1)) if m else None


def _as_entered(value) -> str:
    """The value as typed: TRUE/FALSE for booleans, floats without binary noise, dates in ISO form."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, float):
        return _trim(round(value, 12)) if math.isfinite(value) else str(value)
    if isinstance(value, _DATES):
        return value.isoformat()
    return str(value).strip()


def number_text(value, number_format=None) -> tuple[str, str | None]:
    """(what a reader of the sheet sees, the other form, kept searchable). People quote the
    displayed 0.418 while the cell holds 0.41830627291, so both forms are written."""
    if isinstance(value, int):                     # booleans included
        return _as_entered(value), None
    x = float(value)
    if not math.isfinite(x):
        return str(x), None
    fmt = str(number_format) if number_format else ""
    if x == int(x) and abs(x) < 1e15:
        whole = int(x)
        return (str(int(x * 100)), f"[exact {whole}]") if "%" in fmt else (str(whole), None)
    exact = _trim(round(x, 12))
    places = _decimals(fmt)
    if "%" in fmt:
        return _trim(round(x * 100, places or 0)), f"[exact {exact}]"
    if places is not None:
        shown = _trim(round(x, places))
        return shown, (f"[exact {exact}]" if shown != exact else None)
    approx = _trim(round(x, 3))
    return exact, (f"(≈{approx})" if approx != exact else None)


def cell_text(cell: SheetCell) -> tuple[str, str | None]:
    """(the text shown for a cell, the other form of its number if there is one)."""
    value = cell.value
    if value is None:
        return ("(no stored result)" if cell.formula else ""), None
    if isinstance(value, (int, float)):            # booleans included
        return number_text(value, cell.fmt)
    if isinstance(value, (str, *_DATES)):
        return _as_entered(value), None
    return str(value), None


def _has_text(cell) -> bool:
    return cell is not None and isinstance(cell.value, str) and bool(cell.value.strip())


def _has_number(cell) -> bool:
    return cell is not None and kind_of(cell.value) == "number"


def _has_value(cell) -> bool:
    return cell is not None and not isinstance(cell.value, (str, type(None)))


def _label_left(grid: dict, row: int, col: int, units: Units) -> str | None:
    """The nearest text to the left, unless a number stands in between. A unit is no label."""
    for c in range(col - 1, col - LEFT_LOOK - 1, -1):
        cell = grid.get((row, c))
        if _has_text(cell) and not units.is_unit(cell.value):
            return cell.value.strip()
        if _has_number(cell):
            return None
    return None


def _title_above(grid: dict, row: int, col: int, units: Units) -> str | None:
    for r in range(row - 1, row - UP_LOOK - 1, -1):
        cell = grid.get((r, col))
        if _has_text(cell) and not units.is_unit(cell.value):
            return cell.value.strip()
    return None


def _surroundings(grid: dict, row: int, col: int, units: Units):
    """(label to the left, unit written to the right, column title above)."""
    label = _label_left(grid, row, col, units)
    beside = grid.get((row, col + 1))
    unit_written = beside.value if beside is not None and units.is_unit(beside.value) else None
    title = _title_above(grid, row, col, units)
    if label is None and title:
        return title, unit_written, None
    return label, unit_written, title


def _jsonable(value):
    """A cell value as JSON holds it: dates as ISO text, infinities and NaN as text."""
    if value is None or isinstance(value, (str, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    return value.isoformat() if isinstance(value, _DATES) else str(value)


def _percent(cell: SheetCell) -> bool:
    return bool(cell.fmt) and "%" in str(cell.fmt) and _has_number(cell)


def cell_records(sheets: list[Sheet], units: Units) -> list[dict]:
    out = []
    for sheet in sheets:
        for (row, col), cell in sorted(sheet.grid.items()):
            if not (cell.formula or cell.kind == "error" or _has_value(cell)):
                continue
            label, unit_written, title = _surroundings(sheet.grid, row, col, units)
            unit, known = units.read(unit_written)
            if not unit and _percent(cell):
                unit, known = "%", True
            out.append({
                "format_version": SHEET_SCHEMA_VERSION, "reader_version": SHEET_READER_VERSION,
                "sheet": sheet.name, "cell": f"{column_name(col)}{row}", "row": row, "col": col,
                "label": label, "column_title": title,
                "value": _jsonable(cell.value), "value_string": _as_entered(cell.value),
                "display": cell_text(cell)[0], "value_class": cell.kind,
                "unit_as_written": None if unit_written is None else str(unit_written).strip(),
                "unit": unit, "unit_norm": known,
                "formula": cell.formula, "number_format": cell.fmt if isinstance(cell.fmt, str) else None,
                "hidden": bool(cell.hidden), "sheet_state": sheet.state,
            })
    return out


def _sheet_lines(sheet: Sheet, units: Units) -> list[str]:
    """One sheet as text: a value per line with label, unit and address, then the loose text."""
    name, grid = sheet.name, sheet.grid
    heading = f"## sheet {name}" + (f"  ({sheet.state})" if sheet.state and sheet.state != "visible" else "")
    top_text = [cell.value.strip() for (r, _c), cell in sorted(grid.items()) if r < 3 and _has_text(cell)]
    if top_text:
        heading += "  · " + " / ".join(top_text[:2])
    lines = [heading] + ([f"   ({sheet.note})"] if sheet.note else [])
    by_row: dict[int, list] = {}
    for (r, c), cell in sorted(grid.items()):
        by_row.setdefault(r, []).append((c, cell))
    for r in range(1, sheet.rows + 1):
        here = [(c, cell) for c, cell in by_row.get(r, ()) if c <= sheet.cols]
        used = set()
        for col, cell in here:
            if not (cell.formula or _has_value(cell)):
                continue
            label, unit_written, title = _surroundings(grid, r, col, units)
            label_at = _nearest_left(grid, r, col, label) if label else None
            if label_at:
                used.add(label_at)
            if unit_written is not None:
                used.add((r, col + 1))
            shown, other = cell_text(cell)
            unit = units.read(unit_written)[0] or ("%" if _percent(cell) else "")
            parts = [f"{label or '(unlabelled)'} = {shown}", unit, other or ""]
            text = " ".join(p for p in parts if p) + f"  @{name}!{column_name(col)}{r}"
            if title and title != label:
                text += f" [column {title}]"
            if cell.hidden:
                text += " [hidden]"
            if cell.formula:
                text += f"  {{{cell.formula}}}"
            lines.append(text)
            used.add((r, col))
        lines += [f"{cell.value.strip()}  @{name}!{column_name(col)}{r}" + (" [hidden]" if cell.hidden else "")
                  for col, cell in here if (r, col) not in used and _has_text(cell)]
    return lines + [""]


def _nearest_left(grid: dict, row: int, col: int, text: str):
    """The address of the closest cell to the left that holds exactly this text."""
    for c in range(col - 1, col - LEFT_LOOK - 1, -1):
        cell = grid.get((row, c))
        if _has_text(cell) and cell.value.strip() == text:
            return row, c
    return None


def _row_limit(rows: int, cols: int, full: bool) -> tuple[int, str]:
    room = max(1, CELLS_PER_SHEET // max(cols, 1))
    if full:
        return (room, f"cut at {room} of {rows} rows (cell limit)") if rows > room else (rows, "")
    if rows * max(cols, 1) > CELLS_PER_SHEET or rows > DATA_TABLE_ROWS:
        first = min(rows, DATA_TABLE_SAMPLE)
        return first, f"a data table: {rows} rows x {cols} columns, the first {first} rows are listed"
    return rows, ""


def _formula_text(value, array_formula: type) -> str | None:
    if isinstance(value, array_formula):
        value = value.text
    return value.strip() if isinstance(value, str) and value.startswith("=") else None


def _hidden_lines(ws) -> tuple[set, set]:
    """Hidden row numbers and hidden column letters of an openpyxl sheet."""
    rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
    cols = {column_name(i) for dim in ws.column_dimensions.values() if dim.hidden and dim.min
            for i in range(dim.min, (dim.max or dim.min) + 1)}
    return rows, cols


def _merged_values(ws_formulas, ws_values, row_limit: int) -> dict:
    """Every cell of a merged range carries the value of its top-left cell."""
    spread = {}
    for area in ws_formulas.merged_cells.ranges:
        top_left = ws_values.cell(area.min_row, area.min_col).value
        if top_left is None:
            top_left = ws_formulas.cell(area.min_row, area.min_col).value
        spread.update({(r, c): top_left for r in range(area.min_row, min(area.max_row, row_limit) + 1)
                       for c in range(area.min_col, area.max_col + 1)})
    return spread


def _xlsx_sheet(ws_f, ws_v, full: bool):
    from openpyxl.worksheet.formula import ArrayFormula

    rows, note = _row_limit(ws_f.max_row or 0, ws_f.max_column or 0, full)
    cols = ws_f.max_column or 0
    rows_hidden, cols_hidden = _hidden_lines(ws_f)
    spread = _merged_values(ws_f, ws_v, rows)
    grid = {}
    pairs = zip(ws_f.iter_rows(min_row=1, max_row=rows, max_col=cols),
                ws_v.iter_rows(min_row=1, max_row=rows, max_col=cols), strict=False)
    for row_f, row_v in pairs:
        for cf, cv in zip(row_f, row_v, strict=False):
            formula = _formula_text(cf.value, ArrayFormula)
            error = cf.data_type == "e"
            value = cv.value
            if value is None and not formula and error:
                value = cf.value                   # an error code is stored in the formula pass only
            if value is None and not formula:
                value = spread.get((cf.row, cf.column))
                if value is None:
                    continue
            if isinstance(value, str) and value.startswith("=") and not formula:
                formula, value = value, None
            grid[(cf.row, cf.column)] = SheetCell(
                value, formula, cf.number_format, cf.row in rows_hidden or column_name(cf.column) in cols_hidden,
                "error" if error and not formula else kind_of(value))
    return grid, rows, cols, note


def _link_target(link) -> str:
    """The file name of an external workbook link, '' when the link names none."""
    target = getattr(getattr(link, "file_link", None), "Target", "") or ""
    return os.path.basename(str(target))


def _open_xlsx(path: str, full: bool = False):
    import openpyxl

    def load(results: bool):
        return openpyxl.load_workbook(win_path(path), data_only=results, keep_vba=False, keep_links=True)

    formulas_book, results_book = load(False), load(True)
    sheets = []
    try:
        for name in formulas_book.sheetnames:
            ws_f, ws_v = formulas_book[name], results_book[name]
            if hasattr(ws_f, "max_row"):           # a chart sheet holds no cells
                grid, rows, cols, note = _xlsx_sheet(ws_f, ws_v, full)
                sheets.append(Sheet(name, grid, rows, cols, ws_f.sheet_state, note))
        links = getattr(formulas_book, "_external_links", None) or []
        linked = sorted({_link_target(link) for link in links} - {""})
        named = sorted(str(n) for n in getattr(formulas_book, "defined_names", {}))
    finally:
        formulas_book.close()
        results_book.close()
    return sheets, linked, named, f"openpyxl {openpyxl.__version__}"


def _xls_format_text(book, ws, r, c):
    try:
        xf = book.xf_list[ws.cell_xf_index(r, c)]
        return book.format_map[xf.format_key].format_str
    except Exception:  # noqa: BLE001 - a missing format only loses the display form
        return None


def _xls_value(xlrd, book, ws, r, c):
    """(value, kind) of one xlrd cell; dates as datetime, booleans and errors marked."""
    value, kind = ws.cell_value(r, c), None
    cell_type = ws.cell_type(r, c)
    if cell_type == xlrd.XL_CELL_DATE:
        try:
            value = xlrd.xldate_as_datetime(value, book.datemode)
        except Exception:  # noqa: BLE001 - an impossible date keeps its serial number
            pass
    elif cell_type == xlrd.XL_CELL_BOOLEAN:
        value, kind = bool(value), "boolean"
    elif cell_type == xlrd.XL_CELL_ERROR:
        value, kind = xlrd.error_text_from_code.get(value, f"#ERROR({value})"), "error"
    return value, kind


def _xls_hidden(info_map: dict) -> set[int]:
    """Row or column numbers (from 1) that xlrd's rowinfo_map or colinfo_map marks hidden."""
    return {n + 1 for n, info in info_map.items() if getattr(info, "hidden", 0)}


def _xls_sheet(xlrd, book, ws, full: bool, with_formats: bool) -> Sheet:
    rows, note = _row_limit(ws.nrows, ws.ncols, full)
    cols = ws.ncols
    rows_hidden = _xls_hidden(getattr(ws, "rowinfo_map", {}))
    cols_hidden = _xls_hidden(getattr(ws, "colinfo_map", {}))
    spread = {}
    for top, bottom, left, right in getattr(ws, "merged_cells", []):
        first = ws.cell_value(top, left)
        spread.update({(r, c): first for r in range(top, min(bottom, rows)) for c in range(left, right)})
    grid = {}
    for r in range(rows):
        for c in range(cols):
            value, kind = _xls_value(xlrd, book, ws, r, c)
            if value in ("", None):
                value = spread.get((r, c))
            if isinstance(value, str):
                value = value.strip()
            if value in ("", None):
                continue
            grid[(r + 1, c + 1)] = SheetCell(value, None, _xls_format_text(book, ws, r, c) if with_formats else None,
                                             (r + 1) in rows_hidden or (c + 1) in cols_hidden, kind or kind_of(value))
    state = _XLS_SHEET_STATES.get(getattr(ws, "visibility", 0), "visible")
    return Sheet(ws.name, grid, rows, cols, state, note)


def open_xls(path: str, full: bool = False, xlsx_copy: str | None = None):
    import xlrd

    try:
        book, with_formats = xlrd.open_workbook(win_path(path), formatting_info=True), True
    except Exception:  # noqa: BLE001 - formatting info is optional and fails on odd files
        book, with_formats = xlrd.open_workbook(win_path(path)), False
    sheets = [_xls_sheet(xlrd, book, ws, full, with_formats) for ws in book.sheets()]
    reader = f"xlrd {xlrd.__version__}"
    remark = "values only: the formulas of an old .xls file need Excel (see `contextmax checkup`)"
    if xlsx_copy:
        try:
            _formulas_from_copy(sheets, xlsx_copy)
            reader, remark = f"xls+xlsx xlrd {xlrd.__version__}", "formulas from the .xlsx copy Excel made"
        except Exception as exc:  # noqa: BLE001 - the values are still good without formulas
            remark = f"formulas from the .xlsx copy failed ({type(exc).__name__}); values only"
    sheets = [s._replace(note=f"{s.note}; {remark}" if s.note else remark) for s in sheets]
    return sheets, [], [], reader


def _formulas_from_copy(sheets: list[Sheet], xlsx_copy: str) -> int:
    """Formulas from Excel's .xlsx copy onto the grid read from the original file. Values are
    never taken from the copy, so a recalculation during conversion cannot change a number."""
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula

    book = openpyxl.load_workbook(win_path(xlsx_copy), keep_vba=False)
    added = 0
    try:
        for sheet in sheets:
            if sheet.name not in book.sheetnames or not sheet.rows or not sheet.cols:
                continue
            for row in book[sheet.name].iter_rows(max_row=sheet.rows, max_col=sheet.cols):
                for source in row:
                    formula = _formula_text(source.value, ArrayFormula)
                    if formula:
                        place = (source.row, source.column)
                        sheet.grid.setdefault(place, SheetCell(None, fmt=source.number_format)).formula = formula
                        added += 1
    finally:
        book.close()
    return added


def _csv_encoding_of(path: str) -> str:
    with open(win_path(path), "rb") as fh:
        start = fh.read(4)
    return "utf-16" if start[:2] in (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE) else "utf-8-sig"


_WHOLE = re.compile(r"[-+]?\d+")
_FRACTION = re.compile(r"[-+]?\d+[,.]\d+")


def _csv_value_of(text: str):
    """Numbers as numbers (a decimal comma accepted), everything else as text."""
    token = text.strip()
    try:
        if _WHOLE.fullmatch(token):
            return int(token)
        if _FRACTION.fullmatch(token):
            return float(token.replace(",", "."))
    except ValueError:
        pass
    return token


def _open_csv(path: str, full: bool = False):
    encoding = _csv_encoding_of(path)

    def text_handle():
        return open(win_path(path), encoding=encoding, errors="replace", newline="")

    with text_handle() as fh:
        sample = fh.read(65536)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    with text_handle() as fh:
        widths = [len(row) for row in csv.reader(fh, dialect)]
    width = max(widths, default=0)
    limit, note = _row_limit(len(widths), width, full)
    grid = {}
    with text_handle() as fh:
        for r, row in enumerate(csv.reader(fh, dialect), 1):
            if r > limit:
                break
            grid.update({(r, c): SheetCell(_csv_value_of(text)) for c, text in enumerate(row, 1) if text.strip()})
    return [Sheet("CSV", grid, limit, width, "visible", note)], [], [], "csv"


def open_book(path: str, full: bool = False, xlsx_copy: str | None = None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _open_xlsx(path, full)
    if ext == ".xls":
        return open_xls(path, full, xlsx_copy)
    if ext == ".csv":
        return _open_csv(path, full)
    raise ValueError(f"no spreadsheet reader for {ext}")


def extract(in_path: str, out_path: str, cells_path: str | None, units: Units,
            aliases: dict | None = None, xlsx_copy: str | None = None) -> dict:
    """Write the readable text and the cell records; return what was read."""
    try:
        sheets, linked, named, reader = open_book(in_path, xlsx_copy=xlsx_copy)
    except Exception as exc:  # noqa: BLE001 - reported, never raised
        return {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    lines = [f"# workbook: {os.path.basename(in_path)}", f"# read with: {reader} | sheets: {len(sheets)}"]
    if aliases:
        words = {w for sheet in sheets for c in sheet.grid.values() if _has_text(c)
                 for w in _LOWER_WORDS.findall(c.value.lower())}
        found = {w: aliases[w] for w in sorted(words) if w in aliases}
        if found:
            lines.append("# also searchable as: " + "; ".join(f"{k} = {v}" for k, v in found.items()))
    if named:
        lines.append("# named ranges: " + ", ".join(named[:40]))
    if linked:
        lines.append("# linked workbooks: " + ", ".join(linked))
    lines.append("")
    for sheet in sheets:
        lines += _sheet_lines(sheet, units)
    write_out(out_path, "\n".join(lines) + "\n")
    report = {"ok": True, "method": reader.split()[0], "pages": len(sheets), "needs_ocr": False,
              "formulas_count": sum(1 for sheet in sheets for c in sheet.grid.values() if c.formula),
              "sheet_names": [s.name for s in sheets], "linked_workbooks": linked,
              "sheet_schema": SHEET_SCHEMA_VERSION, "sheet_reader_version": SHEET_READER_VERSION}
    if cells_path:
        records = cell_records(sheets, units)
        digest = hashlib.sha256()
        Path(os.path.abspath(cells_path)).parent.mkdir(parents=True, exist_ok=True)
        with open(win_path(cells_path), "wb") as fh:
            for rec in records:
                data = (json.dumps(rec, ensure_ascii=False, sort_keys=True) + "\n").encode("utf-8")
                fh.write(data)
                digest.update(data)
        report.update({"records_count": len(records), "cells_sha256": digest.hexdigest()})
    return report


def cells(path: str, sheet: str | None = None, xlsx_copy: str | None = None) -> dict:
    """Every cell of a workbook with its value, formula and address, for live reads and traces."""
    sheets, linked, named, reader = open_book(path, full=True, xlsx_copy=xlsx_copy)
    rows = [{"sheet": s.name, "addr": f"{column_name(c)}{r}", "row": r, "col": c,
             "value": _jsonable(cell.value), "formula": cell.formula, "hidden": cell.hidden}
            for s in sheets if not sheet or s.name == sheet
            for (r, c), cell in sorted(s.grid.items()) if cell.value is not None or cell.formula]
    return {"engine": reader, "sheets": [s.name for s in sheets], "linked_workbooks": linked,
            "defined_names": named, "cells": rows}
