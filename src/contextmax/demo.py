"""A small made-up project that exercises everything: `contextmax demo <folder>` writes it.

Firmware in C with modules that call each other, a macro standing for a function, calls
through function pointers, a stub that duplicates a real function's name, tests, an FPGA in
VHDL with a testbench and an orphan entity, a Python tool and a MATLAB model. Documents:
requirements with an older revision, a design that names code and cites the requirements,
a budget workbook with formulas and one crossed row, a test plan that contradicts a value,
and a PDF test report. Every byte is written here, so the demo is the same on every machine.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

FILES = {
    # == firmware (C) ==
    "firmware/acquire/inc/sensor.h": """#ifndef SENSOR_H
#define SENSOR_H
#define SENSOR_MAX_COUNTS 4095
#define SENSOR_SAMPLE_HZ 2000
int read_sensor(int channel);
int calibrate_sensor(int channel, int offset);
#endif
""",
    "firmware/acquire/src/sensor.c": """#include "sensor.h"
#include "hal_io.h"

/* Read one channel and remove the calibration offset. */
int read_sensor(int channel)
{
    int raw = hal_read_adc(channel);
    return sensor_filter(raw) - calibrate_sensor(channel, 0);
}

/* Offset measured at power-up for this channel. */
int calibrate_sensor(int channel, int offset)
{
    return offset + hal_read_adc(channel) / 16;
}

static int sensor_filter(int raw)
{
    return raw > SENSOR_MAX_COUNTS ? SENSOR_MAX_COUNTS : raw;
}
""",
    "firmware/control/inc/pid.h": """#ifndef PID_H
#define PID_H
#define PID_SETTLE_TIME_MS 12
#define PID_OUTPUT_LIMIT 1000
typedef struct { int kp; int ki; int sum; } pid_t;
int pid_step(pid_t *p, int err);
void pid_reset(pid_t *p);
#endif
""",
    "firmware/control/src/pid.c": """#include "pid.h"

/* One step of the PI controller; the output is clamped to the actuator range. */
int pid_step(pid_t *p, int err)
{
    p->sum += err;
    return clamp_output(p->kp * err + p->ki * p->sum);
}

void pid_reset(pid_t *p)
{
    p->sum = 0;
}

/* Limit to what the PWM stage can deliver. */
int clamp_output(int u)
{
    if (u > PID_OUTPUT_LIMIT) return PID_OUTPUT_LIMIT;
    if (u < -PID_OUTPUT_LIMIT) return -PID_OUTPUT_LIMIT;
    return u;
}
""",
    "firmware/control/src/tune.c": """#include "pid.h"

/* Adjust gains from a step response. */
void tune_gains(pid_t *p, int overshoot)
{
    pid_reset(p);
    p->kp = overshoot > 10 ? p->kp - 1 : p->kp + 1;
    pid_step(p, 0);
}
""",
    "firmware/control/src/loop.c": """#include "pid.h"
#include "sensor.h"
#include "log.h"

typedef struct { void (*start)(void); void (*stop)(void); } ops_t;
static pid_t loop_pid;

/* The control loop: read, control, drive, log. Runs from the 1 kHz timer. */
void run_loop(const ops_t *ops)
{
    int e = 2048 - read_sensor(0);
    int u = pid_step(&loop_pid, e);
    hal_write_pwm(0, u);
    ops->start();
    LOG("loop");
}

void set_mode(int mode)
{
    if (mode == 0) pid_reset(&loop_pid);
    write_log("mode");
}
""",
    "firmware/report/inc/log.h": """#ifndef LOG_H
#define LOG_H
#define LOG(msg) write_log(msg)
void write_log(const char *msg);
#endif
""",
    "firmware/report/src/log.c": """#include "log.h"

/* Append one line to the flight recorder. */
void write_log(const char *msg)
{
    format_entry(msg);
    can_send(0x701, msg);
}

static void format_entry(const char *msg)
{
    (void)msg;
}
""",
    "firmware/hal/inc/hal_io.h": """#ifndef HAL_IO_H
#define HAL_IO_H
int hal_read_adc(int channel);
void hal_write_pwm(int channel, int duty);
#endif
""",
    "firmware/hal/src/hal_io.c": """#include "hal_io.h"

/* Read the ADC register of one channel. */
int hal_read_adc(int channel)
{
    return *(volatile int *)(0x40012000 + 4 * channel);
}

void hal_write_pwm(int channel, int duty)
{
    *(volatile int *)(0x40013000 + 4 * channel) = duty;
}
""",
    "firmware/comms/src/can.c": """/* CAN transport. Handlers are looked up in a table at run time. */
typedef void (*handler_t)(const char *);
static handler_t handlers[4];

void can_send(int id, const char *data)
{
    (void)id; (void)data;
}

void can_receive(int id, const char *data)
{
    handlers[id & 3](data);
    write_log(data);
}
""",
    "firmware/test/test_pid.c": """#include "pid.h"

void test_pid_step(void)
{
    pid_t p = {1, 0, 0};
    pid_step(&p, 5);
    clamp_output(5000);
}
""",
    "firmware/stub/hal_stub.c": """/* Stand-in for the hardware when the code runs on a PC. */
int hal_read_adc(int channel)
{
    return 2000 + channel;
}
""",
    # == FPGA (VHDL) ==
    "logic/top/rtl/top.vhd": """library ieee;
use ieee.std_logic_1164.all;
-- Top of the motion FPGA: two encoders, one PWM generator, one timer.
entity motion_top is
  port ( clk : in std_logic; enc_a : in std_logic; pwm_out : out std_logic );
end entity;
architecture rtl of motion_top is
begin
  enc_x : entity work.quad_decoder port map ( clk => clk, a => enc_a );
  enc_y : entity work.quad_decoder port map ( clk => clk, a => enc_a );
  pwm0 : entity work.pwm_gen port map ( clk => clk, q => pwm_out );
  tick : entity work.sys_timer port map ( clk => clk );
end architecture;
""",
    "logic/encoder/rtl/quad_decoder.vhd": """library ieee;
use ieee.std_logic_1164.all;
entity quad_decoder is
  port ( clk : in std_logic; a : in std_logic; count : out std_logic );
end entity;
architecture rtl of quad_decoder is
begin
  process (clk) begin end process;
end architecture;
""",
    "logic/pwm/rtl/pwm_gen.vhd": """library ieee;
use ieee.std_logic_1164.all;
entity pwm_gen is
  generic ( WIDTH : integer := 10 );
  port ( clk : in std_logic; q : out std_logic );
end entity;
architecture rtl of pwm_gen is
begin
  cnt : entity work.sys_timer port map ( clk => clk );
end architecture;
""",
    "logic/timer/rtl/sys_timer.vhd": """library ieee;
use ieee.std_logic_1164.all;
entity sys_timer is
  port ( clk : in std_logic );
end entity;
architecture rtl of sys_timer is
begin
end architecture;
""",
    "logic/legacy/rtl/old_filter.vhd": """library ieee;
use ieee.std_logic_1164.all;
entity old_filter is
  port ( clk : in std_logic );
end entity;
architecture rtl of old_filter is
begin
end architecture;
""",
    "logic/top/tb/motion_top_tb.vhd": """library ieee;
use ieee.std_logic_1164.all;
entity motion_top_tb is
end entity;
architecture sim of motion_top_tb is
begin
  dut : entity work.motion_top port map ( clk => '0', enc_a => '0', pwm_out => open );
end architecture;
""",
    # == tools (Python) ==
    "tools/calibrate.py": '''"""Bench tool: compute the sensor offsets from a capture file."""
import csv


def load_capture(path):
    """Read one capture as a list of integers."""
    with open(path) as fh:
        return [int(r[0]) for r in csv.reader(fh)]


def offset_of(samples):
    """The calibration offset: the mean of the first hundred samples."""
    head = samples[:100]
    return sum(head) // max(1, len(head))


def main(path):
    print(offset_of(load_capture(path)))
''',
    # == models (MATLAB) ==
    "models/pid_model.m": """function y = pid_model(t, kp, ki)
% Closed-loop step response of the PI controller around the plant.
G = plant_tf(0.02);
y = step_response(G, kp, ki, t);
end

function y = step_response(G, kp, ki, t)
% Simulated response, one value per time step.
y = zeros(size(t));
for k = 2:numel(t)
    y(k) = y(k-1) + G * (kp + ki) * (1 - y(k-1));
end
end
""",
    "models/plant_tf.m": """function G = plant_tf(tau)
% First-order plant gain for a time constant tau (seconds).
G = 1 / (1 + tau * 1000);
end
""",
    # == documents ==
    "docs/Verification/Test plan.md": """# Test plan for the control loop

Document DOC-3001. Verifies DOC-1001 and the Detailed design.

## 1 Scope

The unit test `test_pid_step` drives `pid_step` with a step error and checks `clamp_output`.

## 2 Acceptance

- settle time: 15 ms
- output limit: 1000 counts

The control loop shall settle within the settle time after a step [REQ-012].
""",
    "docs/Project/Project notes.md": """# Project notes

Kick-off held. The team agreed to keep the System requirements under change control.
Next review in the spring.
""",
}


def _docx(title: str, paragraphs: list[tuple[str, str]]) -> bytes:
    """A minimal Word file: a title, headings with real heading styles, plain paragraphs."""
    body = []
    for style, text in paragraphs:
        esc = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        body.append(f"<w:p>{ppr}<w:r><w:t xml:space=\"preserve\">{esc}</w:t></w:r></w:p>")
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                + "".join(body) + "</w:body></w:document>")
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
              '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
              '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/></w:style></w:styles>')
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            f'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{title}</dc:title></cp:coreProperties>')
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
              '<Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
              '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/></Relationships>')
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
        for name, data in (("[Content_Types].xml", ctypes), ("_rels/.rels", rels), ("word/document.xml", document),
                           ("word/styles.xml", styles), ("docProps/core.xml", core)):
            info = zipfile.ZipInfo(name, date_time=(2024, 1, 1, 0, 0, 0))
            info.create_system = 0
            z.writestr(info, data)
    return buf.getvalue()


def _pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with a text layer."""
    def pdf_string(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = "BT /F1 11 Tf 72 760 Td 14 TL " + " ".join(f"({pdf_string(t)}) Tj T*" for t in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            "<< /Title (Test report 3: control loop) >>"]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info 6 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def _workbooks(root: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    rows = [("Control loop error budget", None, None), (None, None, None),
            ("Item", "Value", "Unit"),
            ("sample period", 0.5, "[ms]"),
            ("filter delay", 4, "[ms]"),
            ("pid steps to settle", 15, "[-]"),
            ("settle time", "=B4*B6+B5*0", "[ms]"),
            ("output limit", 1000, "[counts]"),
            ("sensor range", 4095, "[counts]"),
            ("label from row 8", "=B9", "[counts]")]
    for r in rows:
        ws.append(list(r))
    ws["A10"] = "=A8"            # the crossed row: label points at row 8, value at row 9
    wb.properties.title = "Error budget"
    _save(wb, root / "docs/Design/Error budget.xlsx")
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "FMEA"
    for r in (("Failure mode", "Severity", "Occurrence", "RPN"), ("sensor stuck", 8, 3, "=B2*C2"), ("pwm saturates", 6, 4, "=B3*C3")):
        ws2.append(list(r))
    _save(wb2, root / "docs/Risk/Software FMEA.xlsx")


def _save(wb, path: Path) -> None:
    import io

    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    wb.save(buf)
    # rewrite the zip with fixed dates so the bytes never change
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        for item in sorted(src.infolist(), key=lambda i: i.filename):
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                import re
                data = re.sub(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>", b"", data)
            info = zipfile.ZipInfo(item.filename, date_time=(2024, 1, 1, 0, 0, 0))
            info.create_system = 0
            z.writestr(info, data)
    path.write_bytes(out.getvalue())


def write(root: str | Path) -> Path:
    root = Path(root)
    for rel, text in FILES.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8", newline="\n")
    reqs = [("Heading1", "1 Introduction"),
            ("", "Document DOC-1001. This document lists what the control loop must do."),
            ("Heading1", "2 Control loop"),
            ("", "The control loop shall settle within the settle time after a step [REQ-012]."),
            ("", "settle time: 12 ms"),
            ("", "The output shall stay within the output limit of 1000 counts [REQ-013]."),
            ("Heading1", "3 Sensor"),
            ("", "The sensor shall be sampled at 2000 Hz [REQ-020].")]
    (root / "docs/Requirements").mkdir(parents=True, exist_ok=True)
    (root / "docs/Requirements/System requirements v1.2.docx").write_bytes(_docx("System requirements for the control loop", reqs))
    (root / "docs/Requirements/old").mkdir(parents=True, exist_ok=True)
    (root / "docs/Requirements/old/System requirements v1.1.docx").write_bytes(
        _docx("System requirements for the control loop", reqs[:5]))
    design = [("Heading1", "1 Introduction"),
              ("", "Document DOC-2001. Implements DOC-1001, the System requirements for the control loop."),
              ("Heading1", "2 Architecture"),
              ("", "run_loop reads the sensor with read_sensor, calls pid_step and drives hal_write_pwm."),
              ("Heading1", "3 Control loop"),
              ("Heading2", "3.1 PID step"),
              ("", "pid_step accumulates the error; clamp_output limits the result to PID_OUTPUT_LIMIT."),
              ("", "settle time: 12 ms"),
              ("Heading2", "3.2 Tuning"),
              ("", "tune_gains adjusts kp from the overshoot and calls pid_reset first."),
              ("Heading1", "4 Interfaces"),
              ("", "The sensor delivers up to SENSOR_MAX_COUNTS counts. See the Error budget.xlsx for timing.")]
    (root / "docs/Design").mkdir(parents=True, exist_ok=True)
    (root / "docs/Design/Detailed design v2.1.docx").write_bytes(_docx("Detailed design of the control loop", design))
    (root / "docs/Design/review").mkdir(parents=True, exist_ok=True)
    (root / "docs/Design/review/Detailed design v2.1 review.docx").write_bytes(_docx("Detailed design of the control loop", design[:6]))
    _workbooks(root)
    (root / "docs/Verification").mkdir(parents=True, exist_ok=True)
    (root / "docs/Verification/Test report 3.pdf").write_bytes(_pdf([
        "Test report 3: control loop", "Document DOC-3002. Tests per DOC-3001, the Test plan.",
        "1 Results", "test_pid_step passed. settle time: 12 ms measured.", "clamp_output held the output limit."]))
    (root / "connections.example.json").write_text(
        '[{"from": "firmware", "to": "logic", "kind": "bus-registers", "label": "PWM and encoder registers",\n'
        '  "evidence": ["hal_write_pwm", "hal_read_adc"]},\n'
        ' {"from": "firmware", "to": "host PC", "kind": "protocol", "label": "CAN service messages",\n'
        '  "evidence": ["can_send", "can_receive"]}]\n', encoding="utf-8")
    # one fixed date on every file, so two copies of the demo are identical in every respect
    import os
    stamp = 1704103200          # 2024-01-01 10:00 UTC
    for p in sorted(root.rglob("*")):
        if p.is_file():
            os.utime(p, (stamp, stamp))
    return root


VOCABULARY = {"doc_number_patterns": [r"DOC-\d{4}"],
              "priorities": [{"label": "1-design", "folder": "docs", "path": r"^design/"}],
              "requirement_id_patterns": [r"\bREQ-\d{3}\b"]}


def create_project(source: str | Path, workspace: str | None = None):
    """Write the demo into `source` and register it as the project 'Demo project'."""
    from contextmax import project, util

    root = write(source)
    folders = [{"path": str(root / "docs"), "holds": "documents"},
               {"path": str(root / "firmware"), "holds": "code"},
               {"path": str(root / "logic"), "holds": "code"},
               {"path": str(root / "tools"), "holds": "code"},
               {"path": str(root / "models"), "holds": "code"}]
    proj = project.create("Demo project", folders, workspace)
    proj.data["vocabulary"] = dict(VOCABULARY)
    proj.save()
    ref = proj.path("notes")
    ref.mkdir(parents=True, exist_ok=True)
    util.save_text(ref / "connections.json", (root / "connections.example.json").read_text(encoding="utf-8"))
    return proj
